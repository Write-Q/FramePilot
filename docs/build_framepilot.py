from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
BASE = Path(__file__).parent
source = Path(r'D:\浏览器文件\FramePilot项目计划书.docx')
doc = Document(source)
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)
sec = doc.sections[0]
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.75)
sec.right_margin = Inches(.75)
sec.header_distance = Inches(.3)
sec.footer_distance = Inches(.3)
sec.different_first_page_header_footer = False
for style_name, size in [('Normal',11),('Title',24),('Heading 1',17),('Heading 2',12),('Header',9),('Footer',9)]:
    style = doc.styles[style_name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string('000000')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    pf = style.paragraph_format
    pf.line_spacing = 1.18
    pf.space_after = Pt(7 if style_name == 'Normal' else 8)
    pf.space_before = Pt(0 if style_name in ['Normal','Title'] else 9)
    pf.widow_control = True
    if style_name.startswith('Heading'):
        pf.keep_with_next = True
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for tag in ('w:pBdr','w:shd'):
            el=ppr.find(qn(tag))
            if el is not None: ppr.remove(el)
for hf in [sec.header,sec.footer]:
    for el in list(hf._element): hf._element.remove(el)
hp=sec.header.add_paragraph('FramePilot  项目计划与验收')
hp.style='Header'
fp=sec.footer.add_paragraph('岗位对齐优化版  ·  2026年9月12日')
fp.style='Footer'
fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('    ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)

def hyperlink(p,text,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    co=OxmlElement('w:color');co.set(qn('w:val'),'245B85');pr.append(co)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def para(text,style=None):
    p=doc.add_paragraph(style=style)
    if re.fullmatch(r'\[[^\]]+\]\([^)]+\)', text):
        p.paragraph_format.keep_with_next=True
    pos=0
    for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
        p.add_run(text[pos:m.start()]);hyperlink(p,m[1],m[2]);pos=m.end()
    p.add_run(text[pos:])
    return p

def table(rows):
    count=len(rows[0]);t=doc.add_table(rows=0,cols=count)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    width=(sec.page_width-sec.left_margin-sec.right_margin)/914400
    fractions=[.22,.36,.42] if count==3 else [1/count]*count
    if count==2: fractions=[.26,.74]
    if rows[0][0]=='学习方向': fractions=[.20,.64,.16]
    if rows[0][0]=='阶段': fractions=[.23,.21,.56]
    if rows[0][0]=='层次': fractions=[.18,.37,.45]
    if rows[0][0]=='用例': fractions=[.26,.36,.38]
    for c,f in zip(t.columns,fractions): c.width=Inches(width*f)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        border=OxmlElement('w:'+tag)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:border.set(qn('w:'+k),v)
        borders.append(border)
    props.append(borders)
    for i,data in enumerate(rows):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,txt) in enumerate(zip(row.cells,data)):
            cell.width=Inches(width*fractions[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EE' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side,value in [('top','85'),('bottom','85'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
            p.paragraph_format.keep_with_next=(i<2)
            r=p.add_run(txt);r.font.size=Pt(10);r.font.bold=(i==0)
            if j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().font.size=Pt(3)

lines=(BASE/'FramePilot项目计划书_v2.md').read_text(encoding='utf-8').splitlines()
i=0;first=True
while i<len(lines):
    s=lines[i].strip();i+=1
    if not s:continue
    if s=='<!-- PAGE -->':continue
    if s.startswith('|'):
        rows=[[v.strip() for v in s.strip('|').split('|')]]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[v.strip() for v in lines[i].strip().strip('|').split('|')];i+=1
            if all(re.fullmatch(r':?-+:?',v) for v in vals):continue
            rows.append(vals)
        table(rows);continue
    if s.startswith('# '):
        para(s[2:],'Title' if first else 'Heading 1');first=False
    elif s.startswith('## '):para(s[3:],'Heading 2')
    else:para(s)
doc.core_properties.title='FramePilot 视频创作 Agent 项目计划书'
doc.core_properties.subject='岗位对齐优化版 版本2.0'
doc.core_properties.author=''
doc.core_properties.comments='基于用户提供的原计划修订；岗位核查日期2026年9月12日。'
out=BASE/'FramePilot项目计划书_v2_岗位对齐优化版.docx'
doc.save(out)
print(out)
print(f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} size={out.stat().st_size}')
